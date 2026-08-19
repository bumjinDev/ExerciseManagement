# 2. jfr_alloc_집계.py
#
# JFR 기록의 할당 표본(jdk.ObjectAllocationSample)을 집계해 엑셀로 저장한다.
# 2026-08-05 세션에서 할당 지점을 판정할 때 쓴 집계의 재작성판이다(원본 스크립트는 보존되지 않았다).
#
# 단위 원칙: 시트의 값은 JMC Memory 화면과 같은 이름, 같은 단위로 적는다.
#   - Alloc Total(GiB)      : JMC 표의 Alloc Total 열과 같은 값, 같은 1024 기준 단위
#   - Total Allocation(%)   : JMC 표의 비중 열과 같은 값(표본 무게 기준)
#   - 표본 수(Samples)와 표본 비중(%) : JMC Stack Trace 판의 Samples, Percentage와 같은 기준(개수)
#   - MB(100만 바이트)      : 카운터 실측 문서(호출 1회당 9,522.1MB)와 대조할 때만 함께 적는다
#
# 입력: jfr print로 뽑은 텍스트 파일 1개. 만드는 명령은 같은 폴더의 사용법 텍스트(2_1)에 있다.
# 출력: 엑셀 파일 1개(시트 5장)와 콘솔 요약 몇 줄.
#   시트 1 "클래스별 집계"  : JMC의 Class 표 + 클래스마다 스택 줄별 표본 수 분해
#   시트 2 "스택 트레이스"  : 클래스마다 호출 사슬 나무 전체. JMC Stack Trace 판을 표로 옮긴 것
#   시트 3 "문장별 집계"    : 세 문장(1230행 / 1278행 / 1280·1281·1284행)의 표본과 무게 비중
#   시트 4 "스레드별 집계"  : main과 워커의 몫, 카운터 실측(순차 7회 몫)과의 대조
#   시트 5 "판정 요약"      : 판정 문장, 2026-08-05 집계 기준값과의 대조, 카운터 실측과의 교차, 주의

import sys
import re
from collections import defaultdict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:
    print("openpyxl이 필요하다. 명령창에서 pip install openpyxl 을 실행한 뒤 다시 돌린다.")
    sys.exit(1)

# jfr print가 무게를 사람이 읽는 단위로 찍는다. 환산 기준은 jfr 도구와 같은 1024 배수다.
BYTES = {"bytes": 1, "kB": 1024, "MB": 1024 * 1024, "GB": 1024 * 1024 * 1024}

ENGINE_PREFIX = "com.exercisemanagement.challenge.service.TeamFormationEngine."

# 두 편차 메서드 안에서 스트림을 만드는 줄들. 이 순서대로 시트의 줄별 열이 된다.
LINE_COLS = [1230, 1278, 1280, 1281, 1284]

# 2026-08-05 세션 집계의 기준값. 재작성판이 같은 기록 파일에서 같은 값을 내는지 대조하는 용도다.
REF_SAMPLES = 11568
REF_TOTAL_MB = 140659
REF_STMT = {"합": 52.3, "평균": 25.2, "분산": 22.5}
REF_CLASS = {"도구": 10.3, "시작": 19.4, "변환": 19.6, "실행분": 50.7, "캡처람다": 1.9}

# 카운터 실측(2026-08-04, getThreadAllocatedBytes): improve() 호출 1회당 총 할당(MB). 교차 확인용.
COUNTER_PER_CALL_MB = 9522.1
CALL_COUNT = 15          # 이 기록 실행의 improve() 호출 수 = 순차 7회 + 병렬 8회
SEQ_CALLS = 7            # 순차 호출 수(전부 main 스레드 실행)
PAR_CALLS = 8            # 병렬 호출 수(대부분 워커 실행)

DEFAULT_OUT = "2_2. jfr_alloc_집계_병렬도6_분할36.xlsx"


def parse(path):
    """jfr print 출력 텍스트를 읽어 표본 목록을 만든다.
    표본 하나 = {cls: 만들어진 클래스 이름, weight: 무게(바이트), thread: 스레드 이름, frames: 스택 줄 목록}"""
    events = []
    cur = None
    in_stack = False
    with open(path, encoding="utf-8", errors="replace") as f:
        for raw in f:
            s = raw.strip()
            if s.startswith("jdk.ObjectAllocationSample"):
                cur = {"cls": "", "weight": 0.0, "thread": "", "frames": []}
                in_stack = False
                continue
            if cur is None:
                continue
            if in_stack:
                if s == "]":
                    in_stack = False
                else:
                    cur["frames"].append(s)
                continue
            if s.startswith("objectClass ="):
                m = re.match(r"objectClass = (.+?)(?: \(classLoader.*)?$", s)
                if m:
                    cur["cls"] = m.group(1).strip()
            elif s.startswith("weight ="):
                m = re.match(r"weight = ([0-9.]+) (bytes|kB|MB|GB)", s)
                if m:
                    cur["weight"] = float(m.group(1)) * BYTES[m.group(2)]
            elif s.startswith("eventThread ="):
                m = re.search(r'"([^"]*)"', s)
                if m:
                    cur["thread"] = m.group(1)
            elif s.startswith("stackTrace ="):
                in_stack = True
            elif s == "}":
                events.append(cur)
                cur = None
    return events


def deviation_line(frames):
    """스택에서 두 편차 메서드의 프레임을 안쪽부터 찾아 그 줄 번호를 돌려준다.
    없으면 None(두 메서드 밖에서 만들어진 표본)."""
    for fr in frames:
        if ENGINE_PREFIX in fr and (".sumDeviation(" in fr or ".distributionDeviation(" in fr):
            m = re.search(r"line: (-?\d+)", fr)
            if m:
                return int(m.group(1))
    return None


def class_group(cls):
    """클래스 이름을 판정에 쓴 묶음으로 나눈다."""
    if cls == "java.util.ArrayList$ArrayListSpliterator":
        return "목록을 훑는 도구"
    if cls == "java.util.stream.ReferencePipeline$Head":
        return "스트림 시작 단계"
    if cls == "java.util.stream.ReferencePipeline$6":
        return "변환 단계"
    if cls.startswith(ENGINE_PREFIX.rstrip(".") + "$$Lambda"):
        return "캡처 람다"
    if (cls == "double[]"
            or cls.startswith("java.util.stream.ReduceOps")
            or cls == "java.util.stream.ReferencePipeline$6$1"
            or cls.startswith("java.util.stream.DoublePipeline$$Lambda")):
        return "합계·평균 실행분"
    return "기타"


def thread_group(name):
    if name == "main":
        return "main"
    if name.startswith("ForkJoinPool"):
        return "워커"
    return "기타 스레드"


def pct(part, total):
    return 100.0 * part / total if total > 0 else 0.0


def gib(w):
    """바이트를 GiB로 바꾼다. JMC의 Alloc Total 표기와 같은 1024 기준이다."""
    return w / (1024.0 ** 3)


def mb(w):
    """바이트를 MB로 바꾼다. 1MB = 100만 바이트. 카운터 실측 문서와 대조할 때만 쓴다."""
    return w / 1e6


def tree_insert(root, frames):
    """호출 사슬 하나를 나무에 넣는다. frames는 만들어진 지점(안쪽)부터 바깥 순서다."""
    node = root
    for fr in frames:
        node = node["ch"].setdefault(fr, {"n": 0, "ch": {}})
        node["n"] += 1


def tree_rows(node, depth, class_total, out):
    """나무를 표본 수 큰 가지부터 깊이 우선으로 내려가며 행 목록을 만든다. JMC 트리와 같은 배치다."""
    for fr, child in sorted(node["ch"].items(), key=lambda kv: -kv[1]["n"]):
        out.append(("    " * depth + fr, child["n"], round(pct(child["n"], class_total), 1)))
        tree_rows(child, depth + 1, class_total, out)


def main():
    if len(sys.argv) not in (2, 3):
        print("사용법: python \"2. jfr_alloc_집계.py\" <jfr print 출력 텍스트> [출력 엑셀 이름]")
        print("텍스트 파일을 만드는 명령은 2_1 사용법 텍스트에 있다.")
        sys.exit(1)

    in_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) == 3 else DEFAULT_OUT

    events = parse(in_path)
    total_w = sum(e["weight"] for e in events)

    # ── 집계 ──────────────────────────────────────────────
    # 클래스별: 표본 수, 무게 합, 줄별 표본 수, 그리고 호출 사슬 나무
    class_agg = {}
    line_w = defaultdict(float)
    line_n = defaultdict(int)
    thr_w = defaultdict(float)
    thr_n = defaultdict(int)

    for e in events:
        ln = deviation_line(e["frames"])
        if ln in LINE_COLS:
            key = ln
        elif ln is None:
            key = "두 메서드 밖"
        else:
            key = "그 밖의 줄"
        line_w[key] += e["weight"]
        line_n[key] += 1

        c = class_agg.setdefault(e["cls"], {"n": 0, "w": 0.0, "line_n": defaultdict(int),
                                            "tree": {"n": 0, "ch": {}}})
        c["n"] += 1
        c["w"] += e["weight"]
        c["line_n"][key] += 1
        tree_insert(c["tree"], e["frames"])

        g = thread_group(e["thread"])
        thr_w[g] += e["weight"]
        thr_n[g] += 1

    stmt_sum_w = line_w.get(1230, 0.0)
    stmt_mean_w = line_w.get(1278, 0.0)
    stmt_var_w = line_w.get(1280, 0.0) + line_w.get(1281, 0.0) + line_w.get(1284, 0.0)

    grp_w = defaultdict(float)
    for cls, c in class_agg.items():
        grp_w[class_group(cls)] += c["w"]
    exec_with_lambda = grp_w.get("합계·평균 실행분", 0.0) + grp_w.get("캡처 람다", 0.0)

    # ── 엑셀 작성 ──────────────────────────────────────────────
    wb = Workbook()
    bold = Font(bold=True)

    def add_header(ws, cols, widths):
        ws.append(cols)
        for cell in ws[1]:
            cell.font = bold
        for col_letter, width in widths.items():
            ws.column_dimensions[col_letter].width = width

    # 시트 1. 클래스별 집계 — JMC Memory 화면의 Class 표와 같은 열 + 스택 줄별 표본 수 분해.
    ws1 = wb.active
    ws1.title = "클래스별 집계"
    add_header(
        ws1,
        ["클래스", "묶음", "Samples(표본 수)", "Alloc Total(GiB)", "Total Allocation(%)",
         "1230행 표본", "1278행 표본", "1280행 표본", "1281행 표본", "1284행 표본",
         "그 밖의 줄 표본", "두 메서드 밖 표본"],
        {"A": 62, "B": 18, "C": 15, "D": 15, "E": 17,
         "F": 11, "G": 11, "H": 11, "I": 11, "J": 11, "K": 13, "L": 15})
    for cls, c in sorted(class_agg.items(), key=lambda kv: -kv[1]["w"]):
        ws1.append([
            cls, class_group(cls), c["n"], round(gib(c["w"]), 2), round(pct(c["w"], total_w), 2),
            c["line_n"].get(1230, 0), c["line_n"].get(1278, 0), c["line_n"].get(1280, 0),
            c["line_n"].get(1281, 0), c["line_n"].get(1284, 0),
            c["line_n"].get("그 밖의 줄", 0), c["line_n"].get("두 메서드 밖", 0)])
    ws1.append([])
    ws1.append(["합계", "", len(events), round(gib(total_w), 2), 100.0,
                line_n.get(1230, 0), line_n.get(1278, 0), line_n.get(1280, 0),
                line_n.get(1281, 0), line_n.get(1284, 0),
                line_n.get("그 밖의 줄", 0), line_n.get("두 메서드 밖", 0)])
    ws1.append([])
    ws1.append(["Alloc Total(GiB)과 Total Allocation(%)은 JMC Memory 화면의 같은 이름 열과 같은 값이다."])
    ws1.append(["줄별 표본 열은 JMC에서 클래스를 클릭했을 때 Stack Trace 판의 Samples가 엔진 줄별로 갈라지던 그 값이다."])

    # 시트 2. 스택 트레이스 — 클래스마다 호출 사슬 나무 전체. JMC Stack Trace 판을 표로 옮긴 것.
    # 프레임은 객체가 만들어진 지점(안쪽)부터 바깥 호출자 순서이고, 아래 계층은 들여쓰기로 구분한다.
    ws2 = wb.create_sheet("스택 트레이스", 1)
    add_header(ws2, ["클래스 / 프레임(들여쓰기 = 한 칸 바깥 호출자)", "Samples(표본 수)", "표본 비중(%)"],
               {"A": 130, "B": 15, "C": 12})
    for cls, c in sorted(class_agg.items(), key=lambda kv: -kv[1]["w"]):
        row_idx = ws2.max_row + 1
        ws2.append([cls, c["n"], 100.0])
        for cell in ws2[row_idx]:
            cell.font = bold
        rows = []
        tree_rows(c["tree"], 0, c["n"], rows)
        for fr, n, p in rows:
            ws2.append([fr, n, p])
        ws2.append([])
    ws2.append(["JMC Stack Trace 판과 같은 구조다. 표본 비중은 그 클래스의 표본 수를 100으로 한 개수 기준이고, JMC의 Percentage와 같은 기준이다."])

    # 시트 3. 문장별 집계 — 판정(어느 문장이 얼마를 만드나)의 근거 표.
    ws3 = wb.create_sheet("문장별 집계")
    add_header(ws3, ["문장", "Samples(표본 수)", "표본 비중(%)", "Alloc Total(GiB)", "무게 비중(%) - 판정에 쓴 값"],
               {"A": 40, "B": 15, "C": 12, "D": 15, "E": 24})
    rows3 = [
        ("합 스트림 (1230행)", line_n.get(1230, 0), stmt_sum_w),
        ("평균 스트림 (1278행)", line_n.get(1278, 0), stmt_mean_w),
        ("분산 스트림 (1280·1281·1284행 합)",
         line_n.get(1280, 0) + line_n.get(1281, 0) + line_n.get(1284, 0), stmt_var_w),
        ("  분산 중 1280행 (stream()과 캡처 람다)", line_n.get(1280, 0), line_w.get(1280, 0.0)),
        ("  분산 중 1281행 (mapToDouble)", line_n.get(1281, 0), line_w.get(1281, 0.0)),
        ("  분산 중 1284행 (average 실행)", line_n.get(1284, 0), line_w.get(1284, 0.0)),
        ("두 메서드 안 그 밖의 줄", line_n.get("그 밖의 줄", 0), line_w.get("그 밖의 줄", 0.0)),
        ("두 메서드 밖 (작업 객체 등)", line_n.get("두 메서드 밖", 0), line_w.get("두 메서드 밖", 0.0)),
        ("전체", len(events), total_w)]
    for name, n, w in rows3:
        ws3.append([name, n, round(pct(n, len(events)), 2),
                    round(gib(w), 2), round(pct(w, total_w), 2)])
    ws3.append([])
    ws3.append(["표본 비중은 개수 기준(JMC Stack Trace의 Percentage와 같은 기준)이고, 무게 비중은 바이트 기준이다."])
    ws3.append(["2026-08-05 판정 값(합 52.3 / 평균 25.2 / 분산 22.5)은 무게 비중이다."])

    # 시트 4. 스레드별 집계 — main과 워커의 몫. 카운터 실측과 대조하는 표라서 MB(100만 바이트)를 함께 적는다.
    ws4 = wb.create_sheet("스레드별 집계")
    add_header(ws4, ["스레드", "Samples(표본 수)", "Alloc Total(GiB)", "무게 비중(%)", "MB(100만 바이트)"],
               {"A": 22, "B": 15, "C": 15, "D": 12, "E": 16})
    for g in ("main", "워커", "기타 스레드"):
        ws4.append([g, thr_n.get(g, 0), round(gib(thr_w.get(g, 0.0)), 2),
                    round(pct(thr_w.get(g, 0.0), total_w), 2), round(mb(thr_w.get(g, 0.0)), 0)])
    ws4.append(["전체", len(events), round(gib(total_w), 2), 100.0, round(mb(total_w), 0)])
    ws4.append([])
    ws4.append(["대조. 순차 호출 몫 계산(MB)", "", "", "", round(SEQ_CALLS * COUNTER_PER_CALL_MB, 1)])
    ws4.append(["  계산식: 순차 7회 × 호출 1회당 카운터 실측 9,522.1MB. main 행의 MB와 견준다"])
    ws4.append(["대조. 병렬 호출 몫 계산(MB)", "", "", "", round(PAR_CALLS * COUNTER_PER_CALL_MB, 1)])
    ws4.append(["  계산식: 병렬 8회 × 9,522.1MB. 병렬 호출에서도 스캔 기준값 계산은 main이 하므로 워커 행보다 조금 크게 나오는 것이 맞다"])
    ws4.append([])
    ws4.append(["주의. 표본 수의 비율은 스레드에 따라 쏠린다. 몫의 비교는 표본 수가 아니라 무게(GiB 또는 MB)로 한다"])

    # 시트 5. 판정 요약 — 재집계와 2026-08-05 집계 기준값의 대조, 카운터 실측과의 교차.
    ws5 = wb.create_sheet("판정 요약")
    add_header(ws5, ["항목", "이 파일 재집계", "2026-08-05 집계 기준값", "설명"],
               {"A": 36, "B": 16, "C": 20, "D": 84})
    ws5.append(["판정", "", "",
                "improve() 호출 1회당 총 할당 9.5GB는 두 편차 메서드가 스트림을 만드는 세 문장에서 나온다"])
    ws5.append(["표본 수", len(events), REF_SAMPLES, "할당 표본(jdk.ObjectAllocationSample)의 개수"])
    ws5.append(["총량(MB, 100만 바이트)", round(mb(total_w), 0), REF_TOTAL_MB,
                "표본 무게의 합. 카운터 실측 문서와 같은 단위. 무게가 소수 첫째 자리로 반올림되어 있어 끝자리가 조금 다를 수 있다"])
    ws5.append(["총량(GiB, JMC 표기)", round(gib(total_w), 1), "",
                "같은 값을 JMC의 Alloc Total 단위로 적은 것"])
    ws5.append(["합 스트림(1230행) 무게 비중(%)", round(pct(stmt_sum_w, total_w), 1), REF_STMT["합"], "sumDeviation의 스트림 한 줄"])
    ws5.append(["평균 스트림(1278행) 무게 비중(%)", round(pct(stmt_mean_w, total_w), 1), REF_STMT["평균"], "distributionDeviation의 평균 스트림"])
    ws5.append(["분산 스트림(1280·1281·1284행) 무게 비중(%)", round(pct(stmt_var_w, total_w), 1), REF_STMT["분산"], "distributionDeviation의 분산 스트림"])
    ws5.append(["목록을 훑는 도구 비중(%)", round(pct(grp_w.get("목록을 훑는 도구", 0.0), total_w), 1), REF_CLASS["도구"], "ArrayList$ArrayListSpliterator"])
    ws5.append(["스트림 시작 단계 비중(%)", round(pct(grp_w.get("스트림 시작 단계", 0.0), total_w), 1), REF_CLASS["시작"], "ReferencePipeline$Head"])
    ws5.append(["변환 단계 비중(%)", round(pct(grp_w.get("변환 단계", 0.0), total_w), 1), REF_CLASS["변환"], "ReferencePipeline$6"])
    ws5.append(["합계·평균 실행분 비중(%)", round(pct(exec_with_lambda, total_w), 1), REF_CLASS["실행분"],
                "double[], ReduceOps 두 개, ReferencePipeline$6$1, DoublePipeline 람다, 캡처 람다의 합"])
    ws5.append(["  그중 캡처 람다 비중(%)", round(pct(grp_w.get("캡처 람다", 0.0), total_w), 1), REF_CLASS["캡처람다"],
                "TeamFormationEngine$$Lambda. mean을 담고 다녀 팀마다 새로 만들어진다. 표본은 전부 1280행"])
    ws5.append(["카운터 실측과의 교차(%)",
                round(100.0 * mb(total_w) / (COUNTER_PER_CALL_MB * CALL_COUNT), 1), 98.5,
                "총량(MB) ÷ (호출 15회 × 호출 1회당 카운터 실측 9,522.1MB). 표본 추정이 실측을 받쳐 주는 확인"])
    ws5.append([])
    ws5.append(["주의 1", "", "", "GiB와 % 값은 JMC와 같은 표본 무게 합산이다. 총량의 확정 값은 카운터 실측(getThreadAllocatedBytes)을 쓴다"])
    ws5.append(["주의 2", "", "", "JFR을 켠 실행이므로 이 실행의 시간 값은 판정에 쓰지 않는다"])
    ws5.append(["주의 3", "", "", "원본은 같은 폴더의 1_1 기록 파일이다. 이 엑셀은 2. jfr_alloc_집계.py가 만든 산출물이다"])

    wb.save(out_path)

    # ── 콘솔 요약 ──────────────────────────────────────────────
    print("=== JFR 할당 표본 집계 ===")
    print(f"표본 수: {len(events):,}개 (기준값 {REF_SAMPLES:,})")
    print(f"총량: {gib(total_w):,.1f}GiB = {mb(total_w):,.0f}MB (기준값 {REF_TOTAL_MB:,}MB)")
    print(f"문장별 무게 비중: 합 {pct(stmt_sum_w, total_w):.1f}% / 평균 {pct(stmt_mean_w, total_w):.1f}%"
          f" / 분산 {pct(stmt_var_w, total_w):.1f}% (기준값 52.3 / 25.2 / 22.5)")
    print(f"엑셀 저장 완료: {out_path}")


if __name__ == "__main__":
    main()
