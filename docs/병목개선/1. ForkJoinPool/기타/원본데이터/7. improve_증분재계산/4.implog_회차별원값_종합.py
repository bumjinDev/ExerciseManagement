# -*- coding: utf-8 -*-
"""상세 로그 엑셀(_provelog가 붙지 않은 .xlsx) 여러 개의 「회차별 원값」 시트를 한 파일로 모으는 스크립트.

implog_to_xlsx_v2.py가 실행 하나마다 만들어 둔 .xlsx에서 「회차별 원값」 시트만 꺼내
실행끼리 나란히 놓고 대조한다. 원본 파일은 열어서 읽기만 한다.

「회차별 원값」은 설정 하나의 측정 5회를 평균 내지 않은 원값이다.
프로브 표(_provelog.xlsx의 「정리」)는 5회 평균이라 회차 사이 흔들림이 보이지 않지만,
이 시트에는 회차 하나하나가 그대로 있어 흔들림을 직접 볼 수 있다.

만드는 시트:
  1) 1. 종합   : 아래 열한 블록. 할당을 앞에 놓았다.
       A. 가져온 실행 목록
       B. 할당 한눈표 — 설정 × 실행, 총 할당과 main·워커·기타 나눔, 순차 대비 배수, 분할 1개당
       C. 총 할당(MB) 회차 원값
       D. main 할당(MB) 회차 원값
       E. 워커 할당(MB) 회차 원값
       F. 기타 할당(MB) 회차 원값
       G. elapsed(ms) 회차 원값 — 설정 하나에 실행 여러 줄이 붙는다
       H. 회차 흔들림 한눈표 — 설정 × 실행, 값은 그 실행 회차들의 최대÷최소
       I. 튄 회차 목록 — 그 설정·그 실행의 최솟값 대비 배수가 큰 것부터
       J. 회차 자리별 경향 — 1회차부터 마지막 회차까지 어느 자리가 느린가
       K. busyCore(개) 회차 원값
  2) 2번째부터 : 가져온 「회차별 원값」 시트 원본을 실행 순서대로 하나씩. 값을 고치지 않는다.

읽는 대상:
  기본은 이 스크립트가 있는 폴더의 .xlsx 중
  이름에 _provelog가 없고 「회차별 원값」 시트가 있는 것 전부다.
  이름이 ~$로 시작하는 엑셀 임시 파일은 건너뛴다.
  --except 로 준 낱말이 파일 이름에 들어 있으면 건너뛴다(기본값 '말단작업분할개수').
  파일을 인자로 직접 주면 그 파일들만 그 순서대로 읽는다.

실행 순서:
  파일 이름의 날짜와 시각(..._20260813_21_21.xlsx)으로 정렬한다.

사용법:
  py -3 "4.implog_회차별원값_종합.py"
  py -3 "4.implog_회차별원값_종합.py" -o "결과.xlsx"
  py -3 "4.implog_회차별원값_종합.py" -d "다른폴더" -o "결과.xlsx"
  py -3 "4.implog_회차별원값_종합.py" --except "말단작업분할개수" --except "19_36"
  py -3 "4.implog_회차별원값_종합.py" 파일1.xlsx 파일2.xlsx -o "결과.xlsx"

필요한 패키지: openpyxl (없으면 pip install openpyxl)
"""
import argparse
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl이 필요하다. pip install openpyxl")

SHEET = "회차별 원값"

FMT_MS = "0.00"
FMT_MB = "0.0000"
FMT_RATIO = "0.00"
FMT_RATIO3 = "0.000"

HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
TITLE_FILL = PatternFill("solid", fgColor="FCE4D6")
SEQ_FILL = PatternFill("solid", fgColor="F2F2F2")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")     # 튄 회차
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")      # 크게 튄 회차
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 튄 회차 기준. 그 설정·그 실행의 5회차 최솟값의 몇 배를 넘으면 튄 것으로 본다.
SPIKE = 1.4
BIG_SPIKE = 2.0

C_IMPL, C_PAR, C_SPLIT, C_RUN = "구현", "병렬도", "분할", "회차"


# ══════════════════════════════════════════════════════════════════════════
# 1. 원본 읽기
# ══════════════════════════════════════════════════════════════════════════
class Run:
    def __init__(self, path):
        self.path = Path(path)
        self.tag = _tag_of(self.path)
        self.header = []
        self.idx = {}
        self.rows = []
        self.raw_grid = []
        self.cond = {}

    def series(self, key, colname):
        """설정 하나의 회차 1~5 값을 차례로 돌려준다. 없으면 None."""
        if colname not in self.idx:
            return [None] * 5
        out = {}
        for r in self.rows:
            if _key_of(r, self.idx) == key:
                run_no = r[self.idx[C_RUN]]
                if isinstance(run_no, (int, float)):
                    out[int(run_no)] = r[self.idx[colname]]
        n = max(out) if out else 5
        return [out.get(i) for i in range(1, max(n, 5) + 1)]


def _key_of(row, idx):
    impl = row[idx[C_IMPL]]
    if impl == "순차":
        return ("순차", None, None)
    return ("병렬", row[idx[C_PAR]], row[idx[C_SPLIT]])


def _tag_of(path):
    m = re.search(r"_(\d{4})(\d{2})(\d{2})_(\d{2})_?(\d{2})(?:\D|$)", path.stem)
    if m:
        return "%s%s_%s%s" % (m.group(2), m.group(3), m.group(4), m.group(5))
    return path.stem[-20:]


def _sort_key_of(path):
    m = re.search(r"_(\d{8})_(\d{2})_?(\d{2})(?:\D|$)", path.stem)
    if m:
        return (0, m.group(1) + m.group(2) + m.group(3), path.name)
    return (1, "", path.name)


def read_run(path):
    run = Run(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError("「%s」 시트가 없다" % SHEET)

    ws = wb[SHEET]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    run.raw_grid = grid

    hrow = None
    for i, r in enumerate(grid):
        if r and r[0] == C_IMPL:
            hrow = i
            break
    if hrow is None:
        wb.close()
        raise ValueError("「%s」 시트에서 머리줄을 못 찾았다" % SHEET)

    head = grid[hrow]
    ncol = 0
    for c in head:
        if c is None or c == "":
            break
        ncol += 1
    run.header = [str(c) for c in head[:ncol]]
    run.idx = {name: i for i, name in enumerate(run.header)}
    for r in grid[hrow + 1:]:
        if not r or r[0] not in ("순차", "병렬"):
            continue
        run.rows.append(list(r[:ncol]))

    if "원본 로그" in wb.sheetnames:
        run.cond = _read_condition(wb["원본 로그"])
    wb.close()

    # 이 엑셀에는 콘솔 머리줄이 없다. 짝이 되는 _provelog.xlsx가 옆에 있으면 거기서 실행 조건을 읽는다.
    if not run.cond:
        mate = run.path.with_name(run.path.stem + "_provelog.xlsx")
        if mate.exists():
            try:
                wb2 = openpyxl.load_workbook(mate, read_only=True, data_only=True)
                if "원본 로그" in wb2.sheetnames:
                    run.cond = _read_condition(wb2["원본 로그"])
                    run.cond["조건 출처"] = mate.name
                wb2.close()
            except Exception:
                pass
    return run


def _read_condition(ws):
    cond = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i > 40:
            break
        line = ""
        for c in row:
            if isinstance(c, str) and len(c) > len(line):
                line = c
        if not line:
            continue
        if "Executing" in line and "시각" not in cond:
            m = re.match(r"\s*(오전|오후)?\s*([\d:]+)", line)
            cond["시각"] = m.group(0).strip() if m else line[:20]
        if "워밍업" in line and "측정" in line:
            m = re.search(r"워밍업\s*(\d+)\s*·\s*측정\s*(\d+)", line)
            if m:
                cond["워밍업"], cond["측정"] = int(m.group(1)), int(m.group(2))
        if line.startswith("스윕 병렬도:"):
            cond["스윕 병렬도"] = line.split(":", 1)[1].strip()
        if line.startswith("스윕 목표 분할 개수:"):
            cond["스윕 분할"] = line.split(":", 1)[1].strip()
    return cond


def all_keys(runs):
    keys = []
    for run in runs:
        for r in run.rows:
            k = _key_of(r, run.idx)
            if k not in keys:
                keys.append(k)
    seq = [k for k in keys if k[0] == "순차"]
    par = sorted([k for k in keys if k[0] == "병렬"],
                 key=lambda k: (k[1] or 0, k[2] or 0))
    return seq + par


def key_label(key):
    return ("순차", "-", "-") if key[0] == "순차" else ("병렬", key[1], key[2])


def nmax(vs):
    v = [x for x in vs if isinstance(x, (int, float))]
    return max(v) if v else None


def nmin(vs):
    v = [x for x in vs if isinstance(x, (int, float))]
    return min(v) if v else None


def navg(vs):
    v = [x for x in vs if isinstance(x, (int, float))]
    return sum(v) / len(v) if v else None


def nmed(vs):
    v = sorted(x for x in vs if isinstance(x, (int, float)))
    if not v:
        return None
    n = len(v)
    return v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2


def _flat(runs, key, colname, nrun):
    """설정 하나의 값을 실행·회차 구분 없이 한 줄로 편다."""
    out = []
    for run in runs:
        out += [v for v in run.series(key, colname)[:nrun] if isinstance(v, (int, float))]
    return out


# ══════════════════════════════════════════════════════════════════════════
# 2. 종합 시트
# ══════════════════════════════════════════════════════════════════════════
def write_summary(ws, runs, keys, nrun):
    tags = [r.tag for r in runs]
    row = 1
    ws.cell(row, 1, "이 시트는 .xlsx 파일들의 「회차별 원값」 시트를 한자리에 모은 것이다."
                    " 값은 원본 그대로이고, 배수와 최대÷최소만 여기서 계산했다.").font = Font(bold=True)
    row += 2

    # ── A. 실행 목록 ──
    row = _title(ws, row, "A. 가져온 실행 목록  (아래 표들의 순서와 같다)")
    _head(ws, row, ["번호", "태그", "실행 시각", "워밍업", "측정", "스윕 병렬도", "스윕 분할 개수",
                    "순차 1회차", "2회차", "3회차", "4회차", "5회차", "순차 최대÷최소", "파일 이름"])
    row += 1
    seqkey = ("순차", None, None)
    for i, run in enumerate(runs, 1):
        s = run.series(seqkey, "elapsed(ms)")[:nrun]
        hi, lo = nmax(s), nmin(s)
        vals = [i, run.tag, run.cond.get("시각", ""), run.cond.get("워밍업", ""),
                run.cond.get("측정", ""), run.cond.get("스윕 병렬도", ""), run.cond.get("스윕 분할", "")]
        vals += s + [(hi / lo) if hi and lo else None, run.path.name]
        _row(ws, row, vals, fmts={j: FMT_MS for j in range(8, 8 + nrun)} | {8 + nrun: FMT_RATIO3})
        row += 1
    row += 2

    order = [k for k in keys if k[0] == "병렬"]

    # ── B. 할당 한눈표 ──
    # 이번 회차의 본 질문이 여기다. 스트림을 걷어 낸 코드가 실제로 객체를 덜 만드는지 본다.
    row = _title(ws, row, "B. 할당 한눈표  —  호출 한 번에 새로 만든 객체가 얼마인가")
    row = _note(ws, row, "실행별 칸은 그 실행 %d회차의 총 할당 평균(MB)이다."
                         " main·워커·기타 평균은 실행 전체를 합쳐 낸 값이고, 셋을 더하면 총 할당이다."
                         " '순차 대비 배수'의 기준은 순차 총 할당의 중앙값이다"
                         "(순차 1회차는 첫 호출이라 값이 크게 나오는 실행이 있어 평균 대신 중앙값을 썼다)."
                         " '분할 1개당'은 총 할당 ÷ 분할 개수다. 이 값이 분할과 무관하게 일정하면"
                         " 할당이 말단 작업 하나당 붙는 몫이라는 뜻이다." % nrun)
    _head(ws, row, ["순번", "구현", "병렬도", "분할"] + ["%s 총(MB)" % t for t in tags]
          + ["총 평균(MB)", "총 중앙값(MB)", "최소", "최대", "최대÷최소",
             "main 평균(MB)", "워커 평균(MB)", "기타 평균(MB)", "워커 몫(%)",
             "순차 대비 배수", "분할 1개당(KB)"])
    row += 1
    seq_base = nmed(_flat(runs, seqkey, "총 할당(MB)", nrun))
    for key in keys:
        lab = key_label(key)
        no = (order.index(key) + 1) if key in order else ""
        per_run = [navg(run.series(key, "총 할당(MB)")[:nrun]) for run in runs]
        mains = _flat(runs, key, "main 할당(MB)", nrun)
        workers = _flat(runs, key, "워커 할당(MB)", nrun)
        others = _flat(runs, key, "기타 할당(MB)", nrun)
        av, lo, hi = navg(per_run), nmin(per_run), nmax(per_run)
        md = nmed(_flat(runs, key, "총 할당(MB)", nrun))
        mv, wv, ov = navg(mains), navg(workers), navg(others)
        tot = (mv or 0) + (wv or 0) + (ov or 0)
        # 배수와 분할 1개당은 중앙값으로 낸다.
        # 첫 호출 한 회차만 유난히 크게 잡히는 일이 있어 평균이 그쪽으로 끌려가기 때문이다.
        vals = [no] + list(lab) + per_run + [av, md, lo, hi, (hi / lo) if hi and lo else None,
                                             mv, wv, ov,
                                             (100 * wv / tot) if tot else None,
                                             (md / seq_base) if md and seq_base else None,
                                             (md * 1000 / key[2]) if md and key[2] else None]
        base = 5 + len(runs)
        fmts = {j: FMT_MB for j in range(5, base + 4)}     # 실행별 + 평균 + 중앙값 + 최소 + 최대
        fmts[base + 4] = FMT_RATIO3                        # 최대÷최소
        for j in range(base + 5, base + 8):                # main·워커·기타
            fmts[j] = FMT_MB
        fmts[base + 8] = FMT_RATIO                         # 워커 몫
        fmts[base + 9] = FMT_RATIO                         # 순차 대비 배수
        fmts[base + 10] = FMT_RATIO3                       # 분할 1개당
        _row(ws, row, vals, fmts=fmts, fill=SEQ_FILL if key[0] == "순차" else None)
        row += 1
    row += 2

    # ── C. 총 할당 회차 원값 ──
    row = _title(ws, row, "C. 총 할당(MB) 회차 원값")
    row = _note(ws, row, "그 호출 한 번에 main과 워커와 기타 Java 스레드가 새로 만든 객체의 합이다."
                         " 입력이 고정이므로 계산 경로가 같으면 회차마다 거의 같아야 한다.")
    row = _matrix(ws, row, runs, keys, "총 할당(MB)", nrun, FMT_MB, mark=False)
    row += 2

    # ── D. main 할당 회차 원값 ──
    row = _title(ws, row, "D. main 할당(MB) 회차 원값")
    row = _note(ws, row, "main 스레드가 새로 만든 객체다. 순차는 전부 여기에 잡히고,"
                         " 병렬은 분할과 병합에 쓰는 객체가 여기에 잡힌다.")
    row = _matrix(ws, row, runs, keys, "main 할당(MB)", nrun, FMT_MB, mark=False)
    row += 2

    # ── E. 워커 할당 회차 원값 ──
    row = _title(ws, row, "E. 워커 할당(MB) 회차 원값")
    row = _note(ws, row, "ForkJoinPool 워커들이 새로 만든 객체의 합이다."
                         " 말단 작업이 teams 복사본을 만드는 몫이 여기에 잡힌다.")
    row = _matrix(ws, row, runs, keys, "워커 할당(MB)", nrun, FMT_MB, mark=False)
    row += 2

    # ── F. 기타 할당 회차 원값 ──
    row = _title(ws, row, "F. 기타 할당(MB) 회차 원값")
    row = _note(ws, row, "main도 워커도 아닌 Java 스레드가 새로 만든 객체다.")
    row = _matrix(ws, row, runs, keys, "기타 할당(MB)", nrun, FMT_MB, mark=False)
    row += 2

    # ── G. elapsed 회차 원값 ──
    row = _title(ws, row, "G. elapsed(ms) 회차 원값  —  설정 하나에 실행 %d줄이 붙는다" % len(runs))
    row = _note(ws, row, "그 설정·그 실행의 %d회차 최솟값을 1.0으로 보고 %.1f배를 넘는 칸은 노란색,"
                         " %.1f배를 넘으면 주황색이다. 오른쪽 '최대÷최소'는 그 회차들 안에서의 폭이다."
                         % (nrun, SPIKE, BIG_SPIKE))
    row = _matrix(ws, row, runs, keys, "elapsed(ms)", nrun, FMT_MS, mark=True)
    row += 2

    # ── H. 회차 흔들림 한눈표 ──
    row = _title(ws, row, "H. 회차 흔들림 한눈표  —  값은 그 실행 %d회차 elapsed의 최대÷최소" % nrun)
    row = _note(ws, row, "1.00이면 회차가 전부 똑같았다는 뜻이다. 이 표에서 큰 칸이 곧 G블록에서 색이 칠해진 자리다.")
    _head(ws, row, ["순번", "구현", "병렬도", "분할"] + tags + ["평균", "최대"])
    row += 1
    for key in keys:
        lab = key_label(key)
        no = (order.index(key) + 1) if key in order else ""
        vals = []
        for run in runs:
            s = run.series(key, "elapsed(ms)")[:nrun]
            hi, lo = nmax(s), nmin(s)
            vals.append((hi / lo) if hi and lo else None)
        _row(ws, row, [no] + list(lab) + vals + [navg(vals), nmax(vals)],
             fmts={j: FMT_RATIO3 for j in range(5, 7 + len(runs))},
             fill=SEQ_FILL if key[0] == "순차" else None)
        for j, v in enumerate(vals):
            if isinstance(v, (int, float)) and v > SPIKE:
                ws.cell(row, 5 + j).fill = BAD_FILL if v > BIG_SPIKE else WARN_FILL
        row += 1
    row += 2

    # ── I. 튄 회차 목록 ──
    row = _title(ws, row, "I. 튄 회차 목록  —  그 설정·그 실행의 %d회차 최솟값 대비 %.1f배를 넘은 elapsed 회차"
                 % (nrun, SPIKE))
    row = _note(ws, row, "배수가 큰 것부터 늘어놓았다. '스윕 순번'은 그 설정이 실행 안에서 몇 번째로 측정됐는지다.")
    _head(ws, row, ["배수", "실행", "스윕 순번", "병렬도", "분할", "회차", "그 회차(ms)",
                    "그 실행 5회차 최소(ms)", "그 실행 5회차 평균(ms)"])
    row += 1
    spikes = []
    for key in keys:
        if key[0] != "병렬":
            continue
        no = order.index(key) + 1
        for run in runs:
            s = run.series(key, "elapsed(ms)")[:nrun]
            lo = nmin(s)
            if not lo:
                continue
            for i, v in enumerate(s, 1):
                if isinstance(v, (int, float)) and v > lo * SPIKE:
                    spikes.append((v / lo, run.tag, no, key[1], key[2], i, v, lo, navg(s)))
    for sp in sorted(spikes, reverse=True):
        _row(ws, row, list(sp),
             fmts={1: FMT_RATIO, 7: FMT_MS, 8: FMT_MS, 9: FMT_MS},
             fill=BAD_FILL if sp[0] > BIG_SPIKE else WARN_FILL)
        row += 1
    if not spikes:
        _row(ws, row, ["튄 회차 없음"])
        row += 1
    row += 1
    _row(ws, row, ["튄 회차 %d개 / 전체 회차 %d개" % (len(spikes), len(order) * len(runs) * nrun)])
    row += 3

    # ── J. 회차 자리별 경향 ──
    row = _title(ws, row, "J. 회차 자리별 경향  —  1회차부터 %d회차까지 어느 자리가 느린가" % nrun)
    row = _note(ws, row, "설정·실행마다 그 5회차의 최솟값을 1.0으로 놓고 각 회차가 몇 배였는지 낸 뒤,"
                         " 같은 자리끼리 평균 냈다. 1회차가 크면 워밍업이 덜 됐다는 뜻이고,"
                         " 뒤 회차가 크면 실행이 진행될수록 나빠졌다는 뜻이다.")
    _head(ws, row, ["묶음"] + ["%d회차" % i for i in range(1, nrun + 1)] + ["튄 회차 수"])
    row += 1
    # 병렬도별로 나눈 뒤, 같은 병렬도 안에서 다시 스윕 앞쪽·뒤쪽으로 나눈다.
    # 병렬도와 자리를 갈라 놓아야 "뒤에서 잴수록 나쁜가"를 따로 볼 수 있다.
    pars = sorted({k[1] for k in order})
    groups = [("전체 병렬 설정", lambda k: k[0] == "병렬")]
    for p in pars:
        same = [k for k in order if k[1] == p]
        half = len(same) / 2
        groups.append(("병렬도 %s 전체" % p, lambda k, p=p: k[0] == "병렬" and k[1] == p))
        groups.append(("  병렬도 %s · 스윕 앞쪽" % p,
                       lambda k, p=p, s=same, h=half: k in s and s.index(k) < h))
        groups.append(("  병렬도 %s · 스윕 뒤쪽" % p,
                       lambda k, p=p, s=same, h=half: k in s and s.index(k) >= h))
    groups.append(("순차", lambda k: k[0] == "순차"))
    for name, pick in groups:
        acc = [[] for _ in range(nrun)]
        cnt = 0
        for key in keys:
            if not pick(key):
                continue
            for run in runs:
                s = run.series(key, "elapsed(ms)")[:nrun]
                lo = nmin(s)
                if not lo:
                    continue
                for i, v in enumerate(s):
                    if isinstance(v, (int, float)):
                        acc[i].append(v / lo)
                        if v > lo * SPIKE:
                            cnt += 1
        _row(ws, row, [name] + [navg(a) for a in acc] + [cnt],
             fmts={j: FMT_RATIO3 for j in range(2, 2 + nrun)})
        row += 1
    row += 2

    # ── K. busyCore ──
    row = _title(ws, row, "K. busyCore(개) 회차 원값")
    row = _note(ws, row, "procCPU ÷ elapsed다. 그 구간에 평균 몇 개 코어가 동시에 돌았나를 뜻한다."
                         " 병렬도에 가까울수록 워커가 놀지 않았다는 뜻이다.")
    row = _matrix(ws, row, runs, keys, "busyCore(개)", nrun, FMT_RATIO, mark=False)

    _fit(ws, [7, 7, 8, 7, 11] + [13] * (nrun + 12))
    ws.freeze_panes = "F1"


def _matrix(ws, row, runs, keys, colname, nrun, fmt, mark):
    """설정 하나에 실행 여러 줄이 붙는 표를 쓴다."""
    order = [k for k in keys if k[0] == "병렬"]
    _head(ws, row, ["순번", "구현", "병렬도", "분할", "실행"]
          + ["%d회차" % i for i in range(1, nrun + 1)]
          + ["평균", "최소", "최대", "최대÷최소"])
    row += 1
    for key in keys:
        lab = key_label(key)
        no = (order.index(key) + 1) if key in order else ""
        for run in runs:
            s = run.series(key, colname)[:nrun]
            hi, lo, av = nmax(s), nmin(s), navg(s)
            vals = [no] + list(lab) + [run.tag] + s + [av, lo, hi,
                                                       (hi / lo) if hi and lo else None]
            fmts = {j: fmt for j in range(6, 6 + nrun + 3)}
            fmts[6 + nrun + 3] = FMT_RATIO3
            _row(ws, row, vals, fmts=fmts,
                 fill=SEQ_FILL if key[0] == "순차" else None)
            if mark and lo:
                for i, v in enumerate(s):
                    if isinstance(v, (int, float)) and v > lo * SPIKE:
                        ws.cell(row, 6 + i).fill = BAD_FILL if v > lo * BIG_SPIKE else WARN_FILL
            row += 1
    return row


# ══════════════════════════════════════════════════════════════════════════
# 3. 원본 시트 옮기기
# ══════════════════════════════════════════════════════════════════════════
def write_raw(ws, run):
    ws.cell(1, 1, "아래는 %s 의 「%s」 시트를 값 그대로 옮긴 것이다. 고친 값은 없다."
            % (run.path.name, SHEET)).font = Font(bold=True)
    for r, line in enumerate(run.raw_grid, 3):
        for c, v in enumerate(line, 1):
            if v is None:
                continue
            cell = ws.cell(r, c, v)
            if isinstance(v, float):
                cell.number_format = FMT_MS if abs(v) >= 1 else FMT_MB
    for r, line in enumerate(run.raw_grid, 3):
        if line and line[0] == C_IMPL:
            for c in range(1, len(line) + 1):
                if ws.cell(r, c).value is not None:
                    ws.cell(r, c).font = Font(bold=True)
                    ws.cell(r, c).fill = HEAD_FILL
            ws.freeze_panes = ws.cell(r + 1, 5)
            break
    _fit(ws, [8, 8, 7, 7] + [13] * 16)


# ══════════════════════════════════════════════════════════════════════════
# 잔손질
# ══════════════════════════════════════════════════════════════════════════
def _title(ws, row, text):
    c = ws.cell(row, 1, text)
    c.font = Font(bold=True, size=12)
    c.fill = TITLE_FILL
    return row + 1


def _note(ws, row, text):
    ws.cell(row, 1, "· " + text).font = Font(size=9, color="595959")
    return row + 1


def _head(ws, row, names):
    for c, name in enumerate(names, 1):
        cell = ws.cell(row, c, name)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _row(ws, row, values, fmts=None, fill=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row, c, v)
        cell.border = BOX
        if fmts and c in fmts:
            cell.number_format = fmts[c]
        if fill is not None:
            cell.fill = fill


def _fit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def safe_sheet_name(name):
    for ch in "[]:*?/\\":
        name = name.replace(ch, "_")
    return name[:31]


# ══════════════════════════════════════════════════════════════════════════
def main():
    ap = argparse.ArgumentParser(description="「회차별 원값」 시트를 한 파일로 모은다.")
    ap.add_argument("files", nargs="*", help="읽을 .xlsx 파일들. 없으면 폴더를 훑는다")
    ap.add_argument("-d", "--dir", default=None, help="훑을 폴더 (기본: 이 스크립트가 있는 폴더)")
    ap.add_argument("-o", "--out", default=None, help="만들 엑셀 경로")
    ap.add_argument("--except", dest="skip", action="append", default=None,
                    help="파일 이름에 이 낱말이 들어가면 건너뛴다 (여러 번 줄 수 있다)")
    args = ap.parse_args()

    skips = args.skip if args.skip is not None else ["말단작업분할개수"]

    if args.files:
        paths = [Path(f) for f in args.files]
    else:
        base = Path(args.dir) if args.dir else Path(__file__).resolve().parent
        paths = [p for p in base.glob("*.xlsx")
                 if not p.name.startswith("~$") and "_provelog" not in p.name]
        paths = [p for p in paths if not any(s in p.name for s in skips)]
        paths.sort(key=_sort_key_of)

    if not paths:
        sys.exit("읽을 .xlsx가 없다.")

    print("후보 파일 %d개" % len(paths))
    runs = []
    for p in paths:
        if not p.exists():
            print("  건너뜀 (없음): %s" % p)
            continue
        try:
            run = read_run(p)
        except Exception as e:
            print("  건너뜀 (%s): %s" % (e, p.name))
            continue
        runs.append(run)
        print("  [%s] %s  — 회차 행 %d개" % (run.tag, p.name, len(run.rows)))
    if not runs:
        sys.exit("읽어 낸 실행이 없다.")

    keys = all_keys(runs)
    nrun = 0
    for run in runs:
        for r in run.rows:
            v = r[run.idx[C_RUN]]
            if isinstance(v, (int, float)):
                nrun = max(nrun, int(v))
    print("설정 %d개 (순차 포함) · 설정당 회차 %d개" % (len(keys), nrun))

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "1. 종합"
    write_summary(ws1, runs, keys, nrun)

    for i, run in enumerate(runs, 2):
        ws = wb.create_sheet(safe_sheet_name("%d. 회차별_%s" % (i, run.tag)))
        write_raw(ws, run)

    out = Path(args.out) if args.out else paths[0].parent / "improve_증분재계산_회차별원값_종합.xlsx"
    wb.save(out)
    print("만든 파일: %s" % out)
    print("시트: %s" % ", ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
